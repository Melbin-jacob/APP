"""Convert a trades-directory spreadsheet into prisma/data/<name>.json for import.

Usage:  python3 scripts/directory/convert_xlsx.py <file.xlsx> prisma/data/windsor-essex-directory.json
Needs:  pip install openpyxl

Expected columns on the "Directory" sheet: #, Business, Trade, Town, Address, Phone, Website, Notes.
The output is committed so the app can import it without spreadsheet tooling;
re-run this script whenever the spreadsheet changes, then restart / re-seed.
"""
import json, re, sys
import openpyxl

SOURCE = "windsor-essex-2026-10"

# Spreadsheet trade -> category slug. Existing categories are reused where they
# mean the same thing; the rest are created by the import.
TRADE_TO_CATEGORY = {
    "Electrical": "electrical", "Plumbing": "plumbing", "Painting": "painting", "Pest Control": "pest-control",
    "Roofing": "roofing", "Moving": "removals", "Automotive Repair & Collision": "motor-garage",
    "Landscaping & Lawn": "gardening",
    "Appliance Repair": "appliance-repair", "Cabinet Making & Millwork": "cabinet-making",
    "Carpentry & Woodworking": "carpentry", "Concrete & Paving": "concrete-paving",
    "Demolition & Junk Removal": "demolition-junk-removal", "Drywall & Plastering": "drywall-plastering",
    "Excavation & Drainage": "excavation-drainage", "Fencing": "fencing", "Flooring & Tile": "flooring-tile",
    "Garage Doors": "garage-doors", "Glass": "glass", "Home Builders": "home-builders", "HVAC": "hvac",
    "Insulation": "insulation", "Locksmith": "locksmith", "Machining & Tool and Die": "machining",
    "Masonry & Stonework": "masonry", "Renovation & General Contractor": "renovation",
    "Siding & Eavestrough": "siding-eavestrough", "Snow Removal & Property Maintenance": "snow-removal",
    "Towing": "towing", "Tree Service": "tree-service", "Waterproofing & Foundation": "waterproofing",
    "Welding & Metal Fabrication": "welding", "Windows & Doors": "windows-doors",
}

# Approximate town centres (lat, lng). Listings are placed at their town; those
# with no town are placed at the centre of Windsor-Essex with a wider radius.
TOWNS = {
    "Windsor": (42.3149, -83.0364), "Tecumseh": (42.3036, -82.8879), "Oldcastle": (42.2290, -82.9620),
    "Belle River": (42.2950, -82.7110), "Lakeshore": (42.2950, -82.7110), "Amherstburg": (42.1009, -83.1080),
    "Maidstone": (42.2120, -82.8840), "Essex": (42.1750, -82.8240), "Harrow": (42.0350, -82.9180),
    "Kingsville": (42.0390, -82.7390), "Leamington": (42.0530, -82.5990), "LaSalle": (42.2210, -83.0560),
    "McGregor": (42.1290, -82.9690), "South Woodslee": (42.2180, -82.7350), "Emeryville": (42.3090, -82.7600),
    "Comber": (42.2380, -82.5490), "Ruthven": (42.0650, -82.6720), "Pelee Island": (41.7800, -82.6600),
    "Pointe-aux-Roches": (42.3070, -82.6380), "Ruscom Station": (42.2690, -82.6400), "Stoney Point": (42.3030, -82.5430),
    "Puce": (42.3020, -82.7810), "Staples": (42.1530, -82.6010), "Cottam": (42.1500, -82.7330),
    "Wheatley": (42.0960, -82.4600), "Woodslee": (42.2180, -82.7350), "Tilbury": (42.2620, -82.4320),
}
REGION = ("Windsor-Essex", 42.2000, -82.8500)
TOWN_FIX = {"Lasalle": "LaSalle", "Mcgregor": "McGregor", "Pointe-Aux-Roches": "Pointe-aux-Roches"}
POSTAL = re.compile(r"\b([ABCEGHJ-NPRSTVXY]\d[ABCEGHJ-NPRSTV-Z])\s?(\d[ABCEGHJ-NPRSTV-Z]\d)\b")


def clean(v):
    v = str(v).strip() if v is not None else ""
    return re.sub(r"\s+", " ", v)


def town_name(raw):
    t = clean(raw).title()
    return TOWN_FIX.get(t, t)


def slugify(s):
    s = re.sub(r"[^a-z0-9]+", "-", s.lower().replace("&", " and ").replace("'", ""))
    return s.strip("-")[:60]


def phone(raw):
    parts = [p.strip() for p in re.split(r"[;/]", clean(raw)) if p.strip()]
    def norm(p):
        digits = re.sub(r"\D", "", p)
        if len(digits) == 11 and digits.startswith("1"):
            digits = digits[1:]
        return f"+1 {digits[:3]}-{digits[3:6]}-{digits[6:]}" if len(digits) == 10 else p
    return [norm(p) for p in parts]


def website(raw):
    w = clean(raw)
    if not w:
        return None
    w = re.sub(r"^https?://", "", w, flags=re.I)
    if " " in w or "." not in w.split("/")[0]:
        return None
    return "https://" + w


def main(src, dest):
    ws = openpyxl.load_workbook(src, read_only=True, data_only=True)["Directory"]
    header = [clean(h) for h in next(ws.iter_rows(max_row=1, values_only=True))]
    col = {name: header.index(name) for name in ["#", "Business", "Trade", "Town", "Address", "Phone", "Website", "Notes"]}
    out, slugs, unknown = [], set(), set()
    for row in ws.iter_rows(min_row=2, values_only=True):
        get = lambda k: row[col[k]] if col[k] < len(row) else None
        name = clean(get("Business"))
        if not name:
            continue
        trade = clean(get("Trade"))
        category = TRADE_TO_CATEGORY.get(trade)
        if not category:
            unknown.add(trade)
            continue
        address, town = clean(get("Address")), town_name(get("Town"))
        # A few "Town" cells hold a street address instead.
        if re.search(r"\d", town):
            address = address or town
            town = ""
        postal = POSTAL.search(address.upper())
        parts = [p.strip() for p in address.split(",")]
        street = parts[0] if parts and re.search(r"\d", parts[0]) else None
        if not town and len(parts) >= 3:
            town = town_name(parts[1])
        known = town in TOWNS
        lat, lng = TOWNS[town] if known else REGION[1:]
        phones = phone(get("Phone"))
        notes = [n for n in [clean(get("Notes"))] if n]
        if len(phones) > 1:
            notes.append("Other numbers listed: " + ", ".join(phones[1:]))
        slug = base = slugify(name) or "business"
        if slug in slugs:
            slug = slugify(f"{name} {town or 'windsor essex'}")
        n = 2
        while slug in slugs:
            slug, n = f"{base}-{n}", n + 1
        slugs.add(slug)
        out.append({
            "ref": f"{SOURCE}:{int(get('#'))}",
            "name": name,
            "slug": slug,
            "category": category,
            "trade": trade,
            "city": town if town else REGION[0],
            "addressLine": street,
            "postcode": f"{postal.group(1)} {postal.group(2)}" if postal else "",
            "latitude": lat,
            "longitude": lng,
            "serviceRadiusKm": 30 if known else 50,
            "phone": phones[0] if phones else "",
            "website": website(get("Website")),
            "notes": "; ".join(notes) or None,
        })
    if unknown:
        sys.exit(f"Unmapped trades (add them to TRADE_TO_CATEGORY): {sorted(unknown)}")
    with open(dest, "w", encoding="utf-8") as f:
        json.dump({"source": SOURCE, "country": "CA", "currency": "CAD", "listings": out}, f, ensure_ascii=False, indent=1)
    print(f"{len(out)} listings -> {dest}")


if __name__ == "__main__":
    main(sys.argv[1], sys.argv[2])
